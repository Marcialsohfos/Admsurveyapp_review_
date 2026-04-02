"""
Module de conversion de données pour DocuMiner AI
"""

import pandas as pd
import os
import json
from typing import Dict, Any

class DataConverter:
    def __init__(self):
        self.output_dir = 'converted'
        os.makedirs(self.output_dir, exist_ok=True)
    
    def convert_data(self, data: Dict[str, Any], output_format: str) -> str:
        """Convertit les données dans le format demandé"""
        
        filename = f"exported_data.{output_format}"
        filepath = os.path.join(self.output_dir, filename)
        
        # Conversion selon le format
        converters = {
            'csv': self._to_csv,
            'xlsx': self._to_excel,
            'json': self._to_json,
            'txt': self._to_txt,
            'docx': self._to_docx,
            'html': self._to_html,
            'dta': self._to_stata,
            'sav': self._to_spss,
            'xml': self._to_xml
        }
        
        converter = converters.get(output_format, self._to_csv)
        
        try:
            converter(data, filepath)
            return filename
        except Exception as e:
            print(f"Erreur conversion {output_format}: {e}")
            # Fallback vers CSV
            fallback_path = filepath.replace(f'.{output_format}', '.csv')
            self._to_csv(data, fallback_path)
            return f"exported_data.csv"
    
    def _to_dataframe(self, data: Dict[str, Any]) -> pd.DataFrame:
        """Convertit les données en DataFrame"""
        
        # Cas d'erreur OCR
        if not data.get('ocr_success', True):
            return pd.DataFrame({
                'Status': ['Erreur OCR'],
                'Message': [data.get('raw_text', 'Aucun texte extrait')[:500]]
            })
        
        # Données budgétaires
        if data.get('type') == 'budget' and data.get('lignes_budgetaires'):
            return pd.DataFrame(data['lignes_budgetaires'])
        
        # Données RH
        if data.get('type') == 'rh_laboratoire' and data.get('personnel_par_grade'):
            return pd.DataFrame(data['personnel_par_grade'])
        
        # Données tabulaires
        if data.get('type') == 'tabular' and data.get('tables'):
            tables = data['tables']
            if tables:
                first_table = tables[0]
                if first_table.get('headers') and first_table.get('rows'):
                    return pd.DataFrame(first_table['rows'], columns=first_table['headers'])
        
        # Données universelles - extraire sections
        sections = data.get('sections', [])
        if sections:
            rows = []
            for section in sections:
                rows.append({
                    'Section': section.get('title', ''),
                    'Contenu': ' '.join(section.get('content', []))[:500]
                })
            return pd.DataFrame(rows)
        
        # Fallback: texte brut
        raw_text = data.get('raw_text', '')
        lines = [line.strip() for line in raw_text.split('\n') if line.strip()][:100]
        
        return pd.DataFrame({
            'Ligne': range(1, len(lines) + 1),
            'Texte': lines
        })
    
    def _to_csv(self, data: Dict[str, Any], filepath: str):
        df = self._to_dataframe(data)
        df.to_csv(filepath, index=False, encoding='utf-8')
    
    def _to_excel(self, data: Dict[str, Any], filepath: str):
        df = self._to_dataframe(data)
        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Données', index=False)
    
    def _to_json(self, data: Dict[str, Any], filepath: str):
        with open(filepath, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2, default=str)
    
    def _to_txt(self, data: Dict[str, Any], filepath: str):
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write("=== DocuMiner AI - Données extraites ===\n\n")
            f.write(f"Type: {data.get('type', 'Inconnu')}\n")
            f.write(f"Type détecté: {data.get('detected_type', 'Inconnu')}\n")
            f.write(f"OCR réussi: {data.get('ocr_success', False)}\n\n")
            f.write("=== CONTENU ===\n")
            f.write(data.get('raw_text', 'Aucun texte')[:10000])
    
    def _to_docx(self, data: Dict[str, Any], filepath: str):
        """Conversion simple en DOCX"""
        try:
            from docx import Document
            doc = Document()
            doc.add_heading('DocuMiner AI - Données extraites', 0)
            doc.add_paragraph(f"Type: {data.get('type', 'Inconnu')}")
            doc.add_paragraph(f"Type détecté: {data.get('detected_type', 'Inconnu')}")
            doc.add_heading('Contenu', level=1)
            
            raw_text = data.get('raw_text', '')
            # Limiter la longueur
            if len(raw_text) > 5000:
                raw_text = raw_text[:5000] + "... (texte tronqué)"
            doc.add_paragraph(raw_text)
            
            doc.save(filepath)
        except ImportError:
            # Fallback vers TXT
            self._to_txt(data, filepath.replace('.docx', '.txt'))
    
    def _to_html(self, data: Dict[str, Any], filepath: str):
        df = self._to_dataframe(data)
        html = df.to_html(index=False)
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(f"<html><head><meta charset='UTF-8'><title>DocuMiner AI</title></head><body>{html}</body></html>")
    
    def _to_stata(self, data: Dict[str, Any], filepath: str):
        try:
            import pyreadstat
            df = self._to_dataframe(data)
            pyreadstat.write_dta(df, filepath)
        except ImportError:
            self._to_csv(data, filepath.replace('.dta', '.csv'))
    
    def _to_spss(self, data: Dict[str, Any], filepath: str):
        try:
            import pyreadstat
            df = self._to_dataframe(data)
            pyreadstat.write_sav(df, filepath)
        except ImportError:
            self._to_csv(data, filepath.replace('.sav', '.csv'))
    
    def _to_xml(self, data: Dict[str, Any], filepath: str):
        df = self._to_dataframe(data)
        xml_content = '<?xml version="1.0" encoding="UTF-8"?>\n<data>\n'
        for _, row in df.iterrows():
            xml_content += '  <record>\n'
            for col, value in row.items():
                if pd.notna(value):
                    xml_content += f'    <{col}>{value}</{col}>\n'
            xml_content += '  </record>\n'
        xml_content += '</data>'
        
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(xml_content)